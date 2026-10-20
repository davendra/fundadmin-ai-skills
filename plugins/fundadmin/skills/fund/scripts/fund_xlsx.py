#!/usr/bin/env python3
"""Render a table spec (JSON) as an Excel workbook or CSV companion file.

The Markdown report stays the canonical output; this writes a spreadsheet of
the same tables for people who work in Excel. It never writes formulas: text
that looks like a formula is stored as text. Outputs are created exclusively
and privately (0600) and never overwrite an existing file or symlink.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import os
import re
import sys
from datetime import date
from pathlib import Path

SPEC_VERSION = '1'
TYPES = {
    'text': '@',
    'number': '#,##0.00',
    'integer': '#,##0',
    'currency': '#,##0.00',
    'percent': '0.00%',
    'date': 'yyyy-mm-dd',
}
MAX_SHEETS = 20
MAX_COLUMNS = 60
MAX_ROWS = 100_000
MAX_TEXT = 32_767  # Excel's per-cell limit
RESERVED_SHEET = 'Notes'
DISCLAIMER = ('AI-generated draft for human review. Not financial, legal, tax or investment advice. '
              'Figures come from the companion Markdown report, which remains the record. '
              'Source documents were processed by the AI host and its model provider; confirm their '
              'confidentiality terms permit that processing.')
FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


class SpecError(ValueError):
    pass


def required(ok, message):
    if not ok:
        raise SpecError(message)


def text(value, label, limit=MAX_TEXT):
    required(isinstance(value, str) and value.strip() and len(value) <= limit, f'{label} must be nonempty text')
    return value


def load_spec(path: Path) -> dict:
    def unique(pairs):
        result = {}
        for key, value in pairs:
            required(key not in result, f'Duplicate JSON key {key!r}')
            result[key] = value
        return result

    def reject_constant(_):
        raise SpecError('Non-finite JSON number')

    raw = path.read_text(encoding='utf-8')
    required(raw.strip(), 'Spec is empty')
    try:
        return json.loads(raw, object_pairs_hook=unique, parse_constant=reject_constant)
    except json.JSONDecodeError as exc:
        raise SpecError(f'Invalid JSON spec: {exc.msg} (line {exc.lineno})') from None


def cell_value(value, kind, where):
    if value is None:
        return None
    if kind == 'text':
        required(isinstance(value, str) and len(value) <= MAX_TEXT, f'{where}: text cell must be a string')
        return value
    if kind == 'date':
        required(isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', value), f'{where}: date must be YYYY-MM-DD')
        try:
            return date.fromisoformat(value)
        except ValueError:
            raise SpecError(f'{where}: invalid date') from None
    required(isinstance(value, (int, float)) and not isinstance(value, bool), f'{where}: {kind} cell must be a number')
    required(math.isfinite(value), f'{where}: number must be finite')
    if kind == 'integer':
        required(float(value).is_integer(), f'{where}: integer cell must be whole')
        return int(value)
    return value


def validate_spec(spec) -> dict:
    required(isinstance(spec, dict), 'Spec must be a JSON object')
    required(spec.get('schema_version') == SPEC_VERSION, f'schema_version must be "{SPEC_VERSION}"')
    allowed = {'schema_version', 'title', 'sheets', 'notes', 'provenance'}
    required(set(spec) <= allowed, f'Unknown spec keys: {sorted(set(spec) - allowed)}')
    text(spec.get('title'), 'title', 200)
    sheets = spec.get('sheets')
    required(isinstance(sheets, list) and 1 <= len(sheets) <= MAX_SHEETS, f'sheets must list 1 to {MAX_SHEETS} sheets')
    names = set()
    clean = []
    for index, sheet in enumerate(sheets, 1):
        where = f'sheet {index}'
        required(isinstance(sheet, dict), f'{where} must be an object')
        required(set(sheet) <= {'name', 'columns', 'rows', 'freeze_header', 'autofilter'}, f'{where}: unknown keys')
        name = text(sheet.get('name'), f'{where} name', 31)
        required(not re.search(r'[\[\]:*?/\\]', name) and name == name.strip(), f'{where}: invalid sheet name')
        required(name.lower() not in names and name.lower() != RESERVED_SHEET.lower(), f'{where}: duplicate or reserved sheet name')
        names.add(name.lower())
        columns = sheet.get('columns')
        required(isinstance(columns, list) and 1 <= len(columns) <= MAX_COLUMNS, f'{where}: columns must list 1 to {MAX_COLUMNS}')
        keys = []
        for column in columns:
            required(isinstance(column, dict) and set(column) <= {'key', 'label', 'type'}, f'{where}: invalid column')
            required(isinstance(column.get('key'), str) and re.fullmatch(r'[a-z][a-z0-9_]{0,63}', column['key']), f'{where}: invalid column key')
            required(column['key'] not in keys, f'{where}: duplicate column key')
            text(column.get('label'), f'{where} column label', 200)
            required(column.get('type', 'text') in TYPES, f'{where}: unknown column type')
            keys.append(column['key'])
        rows = sheet.get('rows')
        required(isinstance(rows, list) and len(rows) <= MAX_ROWS, f'{where}: rows must be a list of at most {MAX_ROWS}')
        typed_rows = []
        for r, row in enumerate(rows, 1):
            required(isinstance(row, dict), f'{where} row {r} must be an object')
            required(set(row) <= set(keys), f'{where} row {r}: keys not in columns')
            typed_rows.append([cell_value(row.get(c['key']), c.get('type', 'text'), f'{where} row {r} {c["key"]}') for c in columns])
        for flag in ('freeze_header', 'autofilter'):
            required(isinstance(sheet.get(flag, True), bool), f'{where}: {flag} must be true or false')
        clean.append({'name': name, 'columns': columns, 'rows': typed_rows,
                      'freeze_header': sheet.get('freeze_header', True), 'autofilter': sheet.get('autofilter', True)})
    notes = spec.get('notes', [])
    required(isinstance(notes, list) and all(isinstance(n, str) and n.strip() and len(n) <= 2000 for n in notes), 'notes must be a list of text')
    provenance = spec.get('provenance', {})
    required(isinstance(provenance, dict) and all(isinstance(k, str) and isinstance(v, str) and len(v) <= 2000 for k, v in provenance.items()), 'provenance must map text to text')
    return {'title': spec['title'], 'sheets': clean, 'notes': notes, 'provenance': provenance}


def safe_csv_text(value):
    """OWASP CSV-injection guard: text that a spreadsheet would evaluate is quoted as text."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    if isinstance(value, date):
        return value.isoformat()
    return '' if value is None else value


def notes_rows(spec):
    rows = [['Title', spec['title']], ['Status', DISCLAIMER]]
    rows += [[key, value] for key, value in spec['provenance'].items()]
    rows += [['Note', note] for note in spec['notes']]
    return rows


def render_csv(spec) -> bytes:
    """One file: each sheet becomes a titled block, then the notes block."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator='\n')
    for sheet in spec['sheets']:
        writer.writerow([f"[Sheet] {sheet['name']}"])
        writer.writerow([safe_csv_text(c['label']) for c in sheet['columns']])
        for row in sheet['rows']:
            writer.writerow([safe_csv_text(v) for v in row])
        writer.writerow([])
    writer.writerow([f'[Sheet] {RESERVED_SHEET}'])
    for row in notes_rows(spec):
        writer.writerow([safe_csv_text(v) for v in row])
    return buffer.getvalue().encode('utf-8-sig')  # BOM so Excel opens UTF-8 correctly


def render_xlsx(spec) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        try:
            from fund_report import install_command  # same directory, shipped together
            fix = install_command()
        except ImportError:  # copied without its siblings
            fix = 'the pinned requirements in a virtual environment'
        raise RuntimeError(f'openpyxl is not installed for {sys.executable}; no workbook was generated. '
                           f'Install the pinned packages ({fix}) or re-run with --format csv.') from None

    def put(ws, row, column, value, number_format=None):
        cell = ws.cell(row=row, column=column, value=value)
        if isinstance(value, str):
            cell.data_type = 's'  # never let text become a formula
        if number_format and value is not None:
            cell.number_format = number_format
        return cell

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='0E1B2E')
    for sheet in spec['sheets']:
        ws = workbook.create_sheet(sheet['name'])
        for c, column in enumerate(sheet['columns'], 1):
            cell = put(ws, 1, c, column['label'])
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        for r, row in enumerate(sheet['rows'], 2):
            for c, (value, column) in enumerate(zip(row, sheet['columns']), 1):
                put(ws, r, c, value, TYPES[column.get('type', 'text')])
        for c, column in enumerate(sheet['columns'], 1):
            longest = max([len(str(column['label']))] + [len(str(row[c - 1])) for row in sheet['rows'][:500] if row[c - 1] is not None])
            ws.column_dimensions[get_column_letter(c)].width = min(60, max(10, longest + 2))
        if sheet['freeze_header']:
            ws.freeze_panes = 'A2'
        if sheet['autofilter'] and sheet['rows']:
            ws.auto_filter.ref = f'A1:{get_column_letter(len(sheet["columns"]))}{len(sheet["rows"]) + 1}'
    notes = workbook.create_sheet(RESERVED_SHEET)
    for r, (label, value) in enumerate(notes_rows(spec), 1):
        put(notes, r, 1, label).font = Font(bold=True)
        put(notes, r, 2, value).alignment = Alignment(wrap_text=True, vertical='top')
    notes.column_dimensions['A'].width = 18
    notes.column_dimensions['B'].width = 100
    workbook.properties.title = spec['title'][:200]
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


NUMBER = re.compile(r'^(?P<neg>\()?(?P<sign>-)?(?P<cur>[$£€])?\s?(?P<sign2>-)?(?P<num>\d{1,3}(?:,\d{3})+|\d+)(?P<frac>\.\d+)?(?P<pct>%)?(?(neg)\))$')


def parse_number(cell: str):
    """A plain, currency or percentage figure as a number; anything else stays text.

    Deliberately narrow: '1,250,000.50', '$1,250.00', '(3,000)', '-12.5%', '8%'.
    '1.2m', 'USD 5', '~3' and dates stay text, so nothing is guessed.
    """
    m = NUMBER.match(cell.replace('\u2212', '-').strip())
    if not m or sum(bool(m[k]) for k in ('neg', 'sign', 'sign2')) > 1:
        return None  # '(-5)' or '-$-5' is ambiguous, so it stays text
    value = float(m['num'].replace(',', '') + (m['frac'] or ''))
    if m['neg'] or m['sign'] or m['sign2']:
        value = -value
    if m['pct']:
        return value / 100, 'percent'
    return value, 'currency' if m['cur'] else 'number'


DELIMITER = re.compile(r'^\s*\|?\s*:?-+:?\s*(?:\|\s*:?-+:?\s*)*\|?\s*$')


def _cells(line: str):
    body = line.strip()
    body = body[1:] if body.startswith('|') else body
    body = body[:-1] if body.endswith('|') and not body.endswith('\\|') else body
    cells = re.split(r'(?<!\\)\|', body)
    return [re.sub(r'\*\*|__|`', '', c).replace('\\|', '|').strip() for c in cells]


def spec_from_markdown(markdown: str, title: str) -> dict:
    """Every pipe table in the report becomes a sheet named after its heading.

    Cells are copied as written. A column becomes numeric only when every
    non-empty cell parses as a number of one kind; otherwise it stays text.
    """
    lines = markdown.splitlines()
    heading, tables, i = 'Table', [], 0
    while i < len(lines):
        line = lines[i]
        if re.match(r'^#{1,6}\s', line):
            heading = re.sub(r'^#{1,6}\s+', '', line).strip() or heading
        if line.lstrip().startswith('|') and i + 1 < len(lines) and DELIMITER.match(lines[i + 1]):
            header = _cells(line)
            rows = []
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                rows.append(_cells(lines[i]))
                i += 1
            tables.append((heading, header, rows))
            continue
        i += 1
    required(tables, 'The report has no Markdown tables to export')
    notes = []
    if len(tables) > MAX_SHEETS:
        notes.append(f'The report has {len(tables)} tables; the first {MAX_SHEETS} are exported.')
        tables = tables[:MAX_SHEETS]
    sheets, used = [], set()
    for heading, header, rows in tables:
        base = re.sub(r'[\[\]:*?/\\]', ' ', heading).strip()[:28].strip() or 'Table'
        if base.lower() == RESERVED_SHEET.lower():
            base = 'Report notes'
        name, n = base, 2
        while name.lower() in used:
            name, n = f'{base[:26]} {n}', n + 1
        used.add(name.lower())
        width = min(len(header), MAX_COLUMNS)
        grid = [(r + [''] * width)[:width] for r in rows]
        columns, typed_rows = [], [dict() for _ in grid]
        for c in range(width):
            values = [row[c] for row in grid]
            parsed = [parse_number(v) if v else None for v in values]
            kinds = {p[1] for p in parsed if p}
            if kinds == {'number', 'currency'}:
                kinds = {'currency'}  # '$0.00' beside '-500.00' is one money column
            # Bare digits (years, IDs, counts like 2026) stay text rather than becoming 2,026.00.
            formatted = any(re.search(r'[.,%$£€()-]', v) for v in values if v)
            numeric = all(p is not None for v, p in zip(values, parsed) if v) and len(kinds) == 1 and formatted
            kind = kinds.pop() if numeric else 'text'
            key = f'c{c + 1}'
            columns.append({'key': key, 'label': header[c] or f'Column {c + 1}', 'type': kind})
            for r, (v, p) in enumerate(zip(values, parsed)):
                typed_rows[r][key] = (p[0] if numeric and p else None) if kind != 'text' else (v or None)
        sheets.append({'name': name, 'columns': columns, 'rows': typed_rows})
    return {'schema_version': SPEC_VERSION, 'title': title[:200], 'sheets': sheets, 'notes': notes + [
        'Tables copied from the Markdown report; figures were not recalculated.']}


def render(spec: dict, fmt: str) -> bytes:
    return render_xlsx(spec) if fmt == 'xlsx' else render_csv(spec)


def write_private(output: Path, payload: bytes) -> Path:
    """Exclusive, private create. Refuses existing files and dangling symlinks."""
    required('..' not in output.parts, 'Parent path traversal is not supported')
    parent = Path(os.path.abspath(output)).parent
    required(not any(p.is_symlink() for p in [parent, *parent.parents]), 'Symlink paths are not supported')
    try:
        fd = os.open(os.path.abspath(output), os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, 'O_NOFOLLOW', 0), 0o600)
    except FileExistsError:
        raise SpecError('Output already exists; existing files and symlinks are never overwritten') from None
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
    except BaseException:
        try:
            os.unlink(output)
        except OSError:
            pass
        raise
    return output


def build(input_path: Path, output_path: Path, fmt: str | None = None, from_markdown: bool = False) -> Path:
    fmt = fmt or output_path.suffix.lower().lstrip('.')
    required(fmt in {'xlsx', 'csv'}, 'Output must end in .xlsx or .csv, or pass --format')
    required(output_path.suffix.lower() == '.' + fmt, f'Output extension must be .{fmt}')
    required(os.path.abspath(input_path) != os.path.abspath(output_path), 'Input and output must be different')
    if from_markdown:
        spec = validate_spec(spec_from_markdown(input_path.read_text(encoding='utf-8'), input_path.stem))
    else:
        spec = validate_spec(load_spec(input_path))
    return write_private(output_path, render(spec, fmt))  # render fully before creating the file


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--input', required=True, help='Table spec JSON, or a Markdown report with --from-markdown')
    parser.add_argument('--from-markdown', action='store_true', help='Build the sheets from the Markdown tables in --input')
    parser.add_argument('--output', required=True, help='New .xlsx or .csv path (never overwritten)')
    parser.add_argument('--format', choices=['xlsx', 'csv'], help='Defaults to the output extension')
    args = parser.parse_args(argv)
    try:
        result = build(Path(args.input), Path(args.output), args.format, args.from_markdown)
    except (SpecError, OSError, RuntimeError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 1
    print(f'{"XLSX" if result.suffix.lower() == ".xlsx" else "CSV"} generated: {result}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
